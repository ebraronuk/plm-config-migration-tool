from src.excel.validator import validate_columns
from src.excel.normalizer import normalize_row
from src.utils.logger import setup_logger

logger = setup_logger()


def read_excel(file_path, data_only=True):
    """
    Excel dosyasını okuyup workbook ve sayfa isimlerini döndürür.
    """
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Excel dosyalarını okuyabilmek için openpyxl gerekiyor") from exc

    try:
        wb = openpyxl.load_workbook(file_path, data_only=data_only)
        logger.info("Workbook yüklendi: %s", file_path)
        return wb, wb.sheetnames
    except Exception as e:
        raise RuntimeError(f"Excel okuma hatası: {e}") from e


def extract_rows(wb, sheet_name=None):
    """
    İlk satırı başlık kabul ederek seçili sayfayı sözlük listesine çevirir.
    """
    sheet = wb[sheet_name] if sheet_name else wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().upper() if h is not None else "" for h in rows[0]]
    parsed_rows = []
    for raw in rows[1:]:
        if not any(raw):
            continue

        row_dict = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            row_dict[header] = raw[idx] if idx < len(raw) else None
        if row_dict:
            parsed_rows.append(row_dict)
    return parsed_rows


def excel_to_records(file_path, sheet_name=None, required_cols=None, normalize=True):
    """
    Excel -> kayıt listesi hattı; isteğe bağlı doğrulama ve normalizasyon içerir.
    """
    wb, _ = read_excel(file_path)
    rows = extract_rows(wb, sheet_name=sheet_name)

    if not rows:
        logger.warning("Sayfa içinde satır bulunamadı: '%s'", sheet_name or wb.active.title)
        return []

    if required_cols:
        missing = validate_columns(rows, required_cols)
        if missing:
            raise ValueError(f"Zorunlu sütunlar eksik: {', '.join(missing)}")

    if normalize:
        rows = [normalize_row(r) for r in rows]

    logger.info("%s kayıt hazırlandı: %s", len(rows), file_path)
    return rows
